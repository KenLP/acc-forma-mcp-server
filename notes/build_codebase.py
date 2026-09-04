# -*- coding: utf-8 -*-
"""
Deck on luyen: toan bo codebase acc-forma-mcp-server.
Nguon noi dung: docs/study/STUDY_toan-bo-codebase.md (commit 234fb05, v0.1.4 + webhooks).

Style: copy nguyen khoi palette + primitives tu FableFramework/kit/assets/deck_style.py
(Autodesk theme + Artifakt fonts + slideLayout7, roundRect pastel nodes, gray arrows,
italic captions, bilingual EN/VI notes).

Chay:  python notes/build_codebase.py
Sua slide nao thi sua script roi chay lai — KHONG sua tay file .pptx.
"""
import os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

_HERE = os.path.dirname(os.path.abspath(__file__))
KIT = r"D:\AIProjects\FableFramework\kit\assets"
SRC = os.path.join(KIT, "deck_theme.pptx")
OUT = os.path.join(_HERE, "ACC_Forma_MCP_Codebase.pptx")

# ---- palette --------------------------------------------------------------
INK="111111"; GRAY="5B5B5B"; WHITE="FFFFFF"; LINE="E8E8E2"
BLUE_F="DCE8FF"; BLUE_L="2F6FED"
MINT_F="DFF7EC"; GREEN="0E9F6E"
AMBER_F="FFF0D4"; AMBER="F59E0B"
RED_F="FFE2E2";  RED="DC2626"
PAPER="F5F5F0";  PAPER_L="D5D5CB"
LAV_F="ECE9FB";  LAV_L="7C5CD6"

# ---- geometry -------------------------------------------------------------
SW, SH = 12192000, 6858000
MX = 419100
RIGHT = 11772900
CW = RIGHT - MX
TY = 2200000
BY = 6120000
CAP_Y = 6300000

ALIGN = {'l':PP_ALIGN.LEFT,'c':PP_ALIGN.CENTER,'r':PP_ALIGN.RIGHT}
ANCH  = {'t':MSO_ANCHOR.TOP,'m':MSO_ANCHOR.MIDDLE,'b':MSO_ANCHOR.BOTTOM}
def C(h): return RGBColor.from_string(h)

prs = Presentation(SRC)
LAYOUT=None
for master in prs.slide_masters:
    for lay in master.slide_layouts:
        if lay.part.partname.endswith("slideLayout7.xml"): LAYOUT=lay
assert LAYOUT is not None
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    rId = sldId.get(qn('r:id'))
    try: prs.part.drop_rel(rId)
    except Exception: pass
    sldIdLst.remove(sldId)

# ---- primitives -----------------------------------------------------------
def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    for t in ('a:buChar','a:buAutoNum','a:buNone'):
        e = pPr.find(qn(t))
        if e is not None: pPr.remove(e)
    pPr.append(pPr.makeelement(qn('a:buNone'), {}))

def _margins(tf,l=64008,r=64008,t=22860,b=22860):
    tf.margin_left=Emu(l); tf.margin_right=Emu(r); tf.margin_top=Emu(t); tf.margin_bottom=Emu(b)

def _runs(tf, lines, align='c', anchor='m'):
    tf.word_wrap=True; tf.vertical_anchor=ANCH[anchor]
    first=True
    for ln in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first=False
        p.alignment=ALIGN[align]; _no_bullet(p)
        if ln.get('space_before') is not None: p.space_before=Pt(ln['space_before'])
        r=p.add_run(); r.text=ln['text']
        f=r.font; f.size=Pt(ln.get('size',12)); f.bold=ln.get('bold',False)
        f.italic=ln.get('italic',False); f.color.rgb=C(ln.get('color',INK))
        if ln.get('font'): f.name=ln['font']

def box(slide, st, x,y,w,h, fill=None, line=None, lw=0.75, radius=None):
    sp=slide.shapes.add_shape(st, Emu(int(x)),Emu(int(y)),Emu(int(w)),Emu(int(h)))
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=C(fill)
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=C(line); sp.line.width=Pt(lw)
    if radius is not None and st==MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sp.adjustments[0]=radius
        except Exception: pass
    sp.shadow.inherit=False
    return sp

def node(slide,x,y,w,h,title,sub=None,fill=BLUE_F,line=LINE,tsize=12.75,ssize=9.5,
         bold=True,tcolor=INK,scolor=GRAY,radius=0.08,align='c',anchor='m'):
    sp=box(slide,MSO_SHAPE.ROUNDED_RECTANGLE,x,y,w,h,fill=fill,line=line,radius=radius)
    _margins(sp.text_frame)
    lines=[{'text':title,'size':tsize,'bold':bold,'color':tcolor}]
    if sub: lines.append({'text':sub,'size':ssize,'bold':False,'color':scolor,'space_before':2})
    _runs(sp.text_frame,lines,align=align,anchor=anchor)
    return sp

def arrow(slide,x,y,w,h,fill=GRAY,direction='right'):
    st={'right':MSO_SHAPE.RIGHT_ARROW,'down':MSO_SHAPE.DOWN_ARROW,
        'left':MSO_SHAPE.LEFT_ARROW,'up':MSO_SHAPE.UP_ARROW}[direction]
    return box(slide,st,x,y,w,h,fill=fill,line=None)

def textbox(slide,x,y,w,h,lines,align='l',anchor='t'):
    tb=slide.shapes.add_textbox(Emu(int(x)),Emu(int(y)),Emu(int(w)),Emu(int(h)))
    _margins(tb.text_frame,l=0,r=0,t=0,b=0)
    _runs(tb.text_frame,lines,align=align,anchor=anchor)
    return tb

def circle(slide,x,y,d,text,fill=BLUE_L,tcolor=WHITE,size=14,line=None):
    sp=box(slide,MSO_SHAPE.OVAL,x,y,d,d,fill=fill,line=line)
    _margins(sp.text_frame,l=0,r=0,t=0,b=0)
    _runs(sp.text_frame,[{'text':text,'size':size,'bold':True,'color':tcolor}],align='c',anchor='m')
    return sp

def codebox(slide,x,y,w,h,code_lines,header=None,hfill=INK):
    yy=y
    if header:
        hh=470000
        hb=box(slide,MSO_SHAPE.ROUNDED_RECTANGLE,x,yy,w,hh,fill=hfill,line=hfill,radius=0.12)
        _margins(hb.text_frame,l=91440,r=91440,t=18288,b=18288)
        _runs(hb.text_frame,[{'text':header,'size':12,'bold':True,'color':WHITE}],align='c',anchor='m')
        yy+=hh+36000; h=h-hh-36000
    cb=box(slide,MSO_SHAPE.ROUNDED_RECTANGLE,x,yy,w,h,fill=WHITE,line=LINE,radius=0.035)
    _margins(cb.text_frame,l=128016,r=80000,t=64008,b=64008)
    _runs(cb.text_frame,[{'text':t,'size':10.5,'color':INK,'font':'Consolas'} for t in code_lines],align='l',anchor='t')
    return cb

def caption(slide,text):
    textbox(slide,MX,CAP_Y,CW,400000,[{'text':text,'size':12.75,'italic':True,'color':GRAY}],align='l',anchor='t')

def band(slide,x,y,w,h,fill=PAPER,line=LINE,radius=0.04):
    return box(slide,MSO_SHAPE.ROUNDED_RECTANGLE,x,y,w,h,fill=fill,line=line,radius=radius)

def head(slide,title,subtitle):
    t=slide.shapes.title; tf=t.text_frame
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.LEFT; _no_bullet(p)
    for r in list(p.runs): r._r.getparent().remove(r._r)
    r=p.add_run(); r.text=title; r.font.size=Pt(28.5); r.font.bold=True; r.font.color.rgb=C(INK)
    sub=None
    for ph in slide.placeholders:
        if ph.placeholder_format.idx==51: sub=ph
    if sub is not None:
        stf=sub.text_frame; stf.clear(); stf.word_wrap=True; stf.vertical_anchor=MSO_ANCHOR.TOP
        p=stf.paragraphs[0]; p.alignment=PP_ALIGN.LEFT; _no_bullet(p)
        r=p.add_run(); r.text=subtitle; r.font.size=Pt(16.5); r.font.color.rgb=C(GRAY)

def setnotes(slide,en,vi):
    ns=slide.notes_slide; tf=ns.notes_text_frame; tf.clear()
    p=tf.paragraphs[0]; r=p.add_run(); r.text="EN: "+en
    tf.add_paragraph()
    p=tf.add_paragraph(); r=p.add_run(); r.text="VI: "+vi

def newslide(title,subtitle):
    slide=prs.slides.add_slide(LAYOUT)
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx not in (0,51):
            ph._element.getparent().remove(ph._element)
    head(slide,title,subtitle)
    return slide

def hrow(slide, items, y, h, x0=MX, x1=RIGHT, gap=260000, tsize=12.75, ssize=9.5, radius=0.08):
    n=len(items); total=x1-x0; w=(total-gap*(n-1))//n
    shapes=[]
    for i,it in enumerate(items):
        x=x0+i*(w+gap)
        fill=it[2] if len(it)>2 else BLUE_F
        line=it[3] if len(it)>3 else LINE
        shapes.append(node(slide,x,y,w,h,it[0],it[1] if len(it)>1 else None,
                           fill=fill,line=line,tsize=tsize,ssize=ssize,radius=radius))
    return shapes,w

def connect(slide, shapes, aw=210000, ah=230000, fill=GRAY):
    for a,b in zip(shapes,shapes[1:]):
        gl=a.left+a.width; gr=b.left
        arrow(slide,gl+((gr-gl)-aw)//2,a.top+a.height//2-ah//2,aw,ah,fill=fill)

def connector(slide,x1,y1,x2,y2,color=GRAY,w=1.0):
    cn=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Emu(int(x1)),Emu(int(y1)),Emu(int(x2)),Emu(int(y2)))
    cn.line.color.rgb=C(color); cn.line.width=Pt(w); cn.shadow.inherit=False
    return cn

def splitbox(slide,x,y,w,h,top,bot,fill=PAPER,line=PAPER_L,tsize=13.5):
    box(slide,MSO_SHAPE.ROUNDED_RECTANGLE,x,y,w,h,fill=fill,line=line,radius=0.06)
    connector(slide,x+130000,y+h//2,x+w-130000,y+h//2,color=line,w=0.75)
    textbox(slide,x,y,w,h//2,[{'text':top,'size':tsize,'bold':True,'color':INK}],align='c',anchor='m')
    textbox(slide,x,y+h//2,w,h//2,[{'text':bot,'size':tsize,'bold':True,'color':INK}],align='c',anchor='m')

# ===========================================================================
# 1 — The Simple Version
# ===========================================================================
s=newslide("acc-forma-mcp-server: Phiên bản đơn giản nhất",
           "Một cái phích cắm giữa AI và dữ liệu công trình — đọc thoải mái, ghi phải xem trước, làm gì cũng có sổ.")
cy=TY+820000
splitbox(s,MX+150000,cy-160000,2250000,1400000,"Claude","MCP client",fill=LAV_F,line=LAV_L)
oc=node(s,SW//2-1750000,cy,3500000,1080000,"acc-forma-mcp-server","46 tools • stdio",
        fill=BLUE_F,line=BLUE_L,tsize=17,ssize=10.5)
splitbox(s,RIGHT-2400000,cy-160000,2250000,1400000,"ACC / Forma","dữ liệu dự án",fill=AMBER_F,line=AMBER)
arrow(s,MX+2500000,cy+430000,320000,240000)
arrow(s,SW//2+1830000,cy+430000,320000,240000)
node(s,SW//2-1750000,cy+1300000,3500000,600000,"safety/ — 10 module","dry-run • allow-list • audit hash chain",
     fill=MINT_F,line=GREEN,tsize=12.5,ssize=9.5)
caption(s,"46 tools, 272 tests / 37 files, ~10.9k dòng src — commit 234fb05 (v0.1.4 + webhooks).")
setnotes(s,
 "This server is a plug between an AI client and Autodesk Construction Cloud. It exposes 46 tools "
 "over stdio. What makes it different from a plain API wrapper is the safety layer: every write is "
 "previewed and approved first, and everything that happened is written to a tamper-evident log.",
 "Server này là cái phích cắm giữa AI client và ACC/Forma, phơi ra 46 tool qua stdio. Khác biệt so với "
 "một API wrapper thường là tầng safety: mọi thao tác ghi đều phải xem trước và được duyệt, và mọi thứ "
 "đã xảy ra đều vào một cuốn sổ không sửa được.")

# ===========================================================================
# 2 — The Problem
# ===========================================================================
s=newslide("Ba nỗi sợ khi cho AI dùng dữ liệu dự án",
           "Ba nỗi sợ này đẻ ra ba trụ cột của codebase — toàn bộ thư mục safety/ tồn tại vì chúng.")
items=[("AI đoán bừa payload","hallucinate issue_subtype_id\nrồi GHI THẲNG vào dự án thật",RED_F,RED),
       ("Không ai biết AI đã làm gì","cuối tuần thấy 30 issue lạ\nkhông truy được ai / khi nào",RED_F,RED),
       ("Bung phạm vi","credential thấy 300 project\nAI chỉ được đụng 2",RED_F,RED)]
sh,_=hrow(s,items,TY+380000,1500000,tsize=14,ssize=10)
ans=[("dry-run + approval token","xem trước, buộc vào payload",MINT_F,GREEN),
     ("audit hash chain","sổ không sửa được",MINT_F,GREEN),
     ("allow-list + scope","chặn từ gốc",MINT_F,GREEN)]
for i,(t,sub,f,l) in enumerate(ans):
    x=sh[i].left
    arrow(s,x+sh[i].width//2-120000,TY+1960000,240000,220000,direction='down')
    node(s,x,TY+2280000,sh[i].width,900000,t,sub,fill=f,line=l,tsize=12.5,ssize=9.5)
caption(s,"Đó là lý do 10 file trong src/safety/ tồn tại — không phải vì thích phức tạp.")
setnotes(s,
 "Three real fears when you let an LLM touch project data: it can hallucinate a payload and write it "
 "for real; nobody can reconstruct what it did; and a credential that sees 300 projects lets it wander. "
 "Each fear produced one pillar of this codebase.",
 "Ba nỗi sợ thật khi cho LLM đụng dữ liệu dự án: nó có thể đoán bừa payload rồi ghi thật; không ai dựng "
 "lại được nó đã làm gì; và credential nhìn thấy 300 project thì nó đi lạc. Mỗi nỗi sợ đẻ ra một trụ cột.")

# ===========================================================================
# 3 — The Big Picture
# ===========================================================================
s=newslide("Bức tranh tổng: bốn tầng, một cửa ải",
           "Mọi tool đều phải đi qua _wrap.ts — tool nào không qua đây là tool không an toàn.")
y=TY+240000; h=760000; gap=150000
rows=[("MCP client  —  stdio JSON-RPC","Claude Desktop / VS Code",LAV_F,LAV_L),
      ("index.ts  +  server.ts","env → auth → đăng ký 46 tool",PAPER,PAPER_L),
      ("tools/_wrap.ts   ★  TOÀN BỘ AN TOÀN Ở ĐÂY","auth → allow-list → readonly → rate → rules → preview → token → execute → audit",BLUE_F,BLUE_L),
      ("tools/<domain>/  —  46 file, mỗi file 1 tool","chỉ khai input schema + gọi apis/ ; KHÔNG tự phòng vệ",PAPER,PAPER_L),
      ("apis/ (10 client)  +  http/client.ts  +  auth/","retry, backoff, 401 refresh, 429 Retry-After",PAPER,PAPER_L)]
for i,(t,sub,f,l) in enumerate(rows):
    yy=y+i*(h+gap)
    ts=13.5 if i==2 else 12.75
    node(s,MX,yy,CW-2600000,h,t,sub,fill=f,line=l,tsize=ts,ssize=9.5)
    if i<len(rows)-1:
        arrow(s,MX+(CW-2600000)//2-110000,yy+h+10000,220000,130000,direction='down')
node(s,RIGHT-2400000,y+1200000,2400000,1500000,"safety/","10 module\n\n~/.acc-forma-mcp/\naudit/*.jsonl",
     fill=MINT_F,line=GREEN,tsize=14,ssize=10)
caption(s,"apis/ chỉ chuyển đổi REST/GraphQL ↔ type TS. Không business logic ở đó.")
setnotes(s,
 "Four layers plus one gate. The client speaks stdio JSON-RPC. index.ts builds the context and "
 "server.ts registers all 46 tools, each wrapped by _wrap.ts. The tools themselves are deliberately "
 "dumb: they declare a schema and call an api client. All safety lives in one file.",
 "Bốn tầng cộng một cửa ải. Client nói stdio JSON-RPC. index.ts dựng context, server.ts đăng ký 46 tool, "
 "mỗi tool đều bọc bằng _wrap.ts. Bản thân các tool cố tình 'ngu': chỉ khai schema và gọi api client. "
 "Toàn bộ an toàn nằm ở một file duy nhất.")

# ===========================================================================
# 4 — Is an MCP Server a Background Service?
# ===========================================================================
s=newslide("MCP server có phải một service chạy nền không?",
           "Không. Nó là một process con, nói chuyện qua stdin/stdout — không port, không HTTP, không daemon.")
left=[("KHÔNG phải","port lắng nghe",RED_F,RED),("KHÔNG phải","HTTP endpoint",RED_F,RED),
      ("KHÔNG phải","daemon chạy nền",RED_F,RED)]
for i,(t,sub,f,l) in enumerate(left):
    node(s,MX,TY+300000+i*820000,3100000,680000,t,sub,fill=f,line=l,tsize=13,ssize=10)
codebox(s,MX+3500000,TY+260000,CW-3500000,1500000,
        ["const transport = new StdioServerTransport();",
         "await server.connect(transport);"],
        header="src/index.ts:92-93")
node(s,MX+3500000,TY+1900000,CW-3500000,1050000,
     "Hệ quả thực tế: muốn test thật phải SPAWN process và ghi JSON vào stdin",
     "không curl được — đây chính là cách các probe trong đợt review verify",
     fill=AMBER_F,line=AMBER,tsize=13,ssize=10)
caption(s,"Hiểu nhầm này phổ biến nhất khi người mới đọc codebase MCP lần đầu.")
setnotes(s,
 "The most common misunderstanding. An MCP server is a child process spawned by the client, talking "
 "JSON-RPC over stdin and stdout. No port, no HTTP. The practical consequence: to test it for real you "
 "must spawn the process and write JSON to its stdin — you cannot curl it.",
 "Hiểu nhầm phổ biến nhất. MCP server là process con do client sinh ra, nói JSON-RPC qua stdin/stdout. "
 "Không port, không HTTP. Hệ quả thực tế: muốn test thật thì phải spawn process và ghi JSON vào stdin, "
 "không curl được.")

# ===========================================================================
# 5 — What is a Tool?
# ===========================================================================
s=newslide("Một 'tool' trong project này là gì?",
           "Một object TS: name + description + inputSchema (zod) + scope + execute() — khai báo, không tự phòng vệ.")
codebox(s,MX,TY+240000,5600000,2600000,
        ["export const createIssueTool:",
         "    MutationToolDef<typeof inputSchema> = {",
         "  name: 'issues_create',",
         "  description: '...',      // LLM doc + Autodesk review",
         "  kind: 'mutation',",
         "  scope: { kind: 'dm' },   // BAT BUOC",
         "  inputSchema,             // zod",
         "  getProjectId: (i) => i.project_id,",
         "  buildPreview: async (...) => {...},",
         "  execute:      async (...) => {...},",
         "};"],
        header="src/tools/_types.ts:82-120")
rx=MX+5900000; rw=CW-5900000
node(s,rx,TY+280000,rw,900000,"description không phải comment","LLM đọc để chọn tool  +  là artifact Autodesk review khi submit",
     fill=AMBER_F,line=AMBER,tsize=13,ssize=9.5)
node(s,rx,TY+1300000,rw,900000,"scope là field BẮT BUỘC","quên khai = LỖI BIÊN DỊCH, không phải bug im lặng",
     fill=MINT_F,line=GREEN,tsize=13,ssize=9.5)
node(s,rx,TY+2320000,rw,520000,"dry_run được TIÊM tự động","server.ts:30-33 — tool không tự khai",
     fill=BLUE_F,line=BLUE_L,tsize=12.5,ssize=9.5)
caption(s,"Tool càng 'ngu' càng tốt: 46 tool × 8 lớp bảo vệ = 368 chỗ có thể quên nếu để tool tự lo.")
setnotes(s,
 "A tool is just a typed object. Two things are easy to miss. First, description is not an internal "
 "comment: the LLM reads it to choose a tool, and Autodesk reviews it at submission. Second, scope is "
 "a required field, so forgetting it is a compile error rather than a silent hole.",
 "Tool chỉ là một object có kiểu. Hai điểm dễ bỏ sót. Một: description không phải comment nội bộ — LLM "
 "đọc nó để chọn tool, và Autodesk review nó khi submit. Hai: scope là field bắt buộc, nên quên khai là "
 "lỗi biên dịch chứ không phải lỗ hổng im lặng.")

# ===========================================================================
# 6 — 46 tools
# ===========================================================================
s=newslide("46 tools, 11 nhóm — 9 là mutation",
           "Đọc chiếm đa số; chỉ 9 tool ghi, và cả 9 đều đi qua hai bước preview + token.")
r1=[("Issues  (12)","list • get • create ✍ • update ✍\ncomment ✍ • pin_element ✍",BLUE_F,BLUE_L),
    ("Data Mgmt  (6)","hubs • projects • folders\nitems • versions",BLUE_F,BLUE_L),
    ("AEC Data Model  (8)","GraphQL: elements, categories,\nparameters, positions",BLUE_F,BLUE_L),
    ("Account Admin  (4)","projects • users • companies",BLUE_F,BLUE_L)]
sh1,_=hrow(s,r1,TY+260000,1250000,tsize=13,ssize=9)
r2=[("Model Derivative (3)","properties • manifest\ntranslation ✍",BLUE_F,BLUE_L),
    ("Reviews  (4)","list • get • create ✍\ntransition ✍",BLUE_F,BLUE_L),
    ("Model Coord.  (2)","modelsets • clashes",BLUE_F,BLUE_L),
    ("Webhooks  (3)","create ✍ • list • delete ✍",AMBER_F,AMBER)]
sh2,_=hrow(s,r2,TY+1640000,1250000,tsize=13,ssize=9)
r3=[("Model Properties (1)","mp_diff_versions — so hai phiên bản",BLUE_F,BLUE_L),
    ("ACC Docs  (1)","viewables cho pin PDF 2D",BLUE_F,BLUE_L),
    ("Meta  (2)","changelog • verify_audit_chain",MINT_F,GREEN)]
hrow(s,r3,TY+3020000,760000,tsize=13,ssize=9)
caption(s,"✍ = mutation. Registry: src/tools/_registry.ts — manifest-sync.spec.ts khoá cho manifest khớp registry.")
setnotes(s,
 "Forty-six tools across eleven domains, of which only nine write. Webhooks is amber because it is the "
 "one mutation that configures ongoing egress: once created, Autodesk keeps POSTing project events to a "
 "third-party URL with no further call from us.",
 "46 tool thuộc 11 nhóm, chỉ 9 tool ghi. Webhooks tô màu amber vì nó là mutation duy nhất cấu hình egress "
 "lâu dài: tạo xong là Autodesk cứ thế POST dữ liệu sự kiện tới URL bên thứ ba, không cần gọi lại từ ta.")

# ===========================================================================
# 7 — What is dry-run + approval token?
# ===========================================================================
s=newslide("Dry-run + approval token là gì?",
           "Thao tác nguy hiểm chia hai bước: bước 1 xem trước và nhận vé, bước 2 đưa vé lại mới được làm thật.")
steps=[("1. Xem trước","gọi tool với dry_run=true",AMBER_F,AMBER),
       ("2. Nhận vé","server cấp approval_token",BLUE_F,BLUE_L),
       ("3. Đối chiếu","token buộc vào payload hash",BLUE_F,BLUE_L),
       ("4. Làm thật","dry_run=false + token",MINT_F,GREEN)]
sh,w=hrow(s,steps,TY+520000,1250000,tsize=13.5,ssize=9.5)
connect(s,sh)
for i,sp in enumerate(sh):
    circle(s,sp.left+w//2-160000,TY+180000,320000,str(i+1))
node(s,MX,TY+2150000,CW,900000,
     "Vé không phải mật khẩu chung — nó bị BUỘC vào đúng payload đã xem",
     "đổi một chữ trong title sau khi preview  →  token vô hiệu  →  ApprovalError",
     fill=RED_F,line=RED,tsize=14,ssize=10.5)
caption(s,"Chống dùng một vé xem hàng A rồi mang đi lấy hàng B.")
setnotes(s,
 "A dangerous operation is split in two: preview first, then confirm. The ticket is not a shared "
 "password — it is cryptographically bound to the exact payload that was previewed. Change one "
 "character of the title between the two calls and the token is void.",
 "Thao tác nguy hiểm chia hai bước: xem trước rồi xác nhận. Cái vé không phải mật khẩu chung — nó bị "
 "buộc bằng mật mã vào đúng payload đã xem. Đổi một ký tự trong title giữa hai lượt là token vô hiệu.")

# ===========================================================================
# 8 — Example, luot 1
# ===========================================================================
s=newslide("Ví dụ — lượt 1: tạo issue (preview)",
           "Không gửi dry_run thì zod điền mặc định true; server chạy hết 6 bước rồi DỪNG, không ghi gì lên ACC.")
codebox(s,MX,TY+220000,5300000,1450000,
        ['{"name":"issues_create","arguments":{',
         '  "project_id":"b.5341a615-...",',
         '  "title":"Thieu cua chong chay",',
         '  "issue_subtype_id":"a1b2..."  }}',
         '// khong co dry_run -> mac dinh TRUE'],
        header="LLM gửi qua stdio")
rx=MX+5600000; rw=CW-5600000; yy=TY+230000; hh=430000; g=90000
pipe=[("0. auth mode","ssa|3lo",PAPER,PAPER_L),
      ("1. allow-list","scope dm → check project",BLUE_F,BLUE_L),
      ("2-3. readonly + rate","≤50/project/giờ",PAPER,PAPER_L),
      ("4. business rules","cục bộ, chưa gọi APS",PAPER,PAPER_L),
      ("5. buildPreview","CÓ gọi APS: kiểm subtype có thật",AMBER_F,AMBER),
      ("6. dry_run → DỪNG","cấp token, audit 'preview'",MINT_F,GREEN)]
for i,(t,sub,f,l) in enumerate(pipe):
    node(s,rx,yy+i*(hh+g),rw,hh,t,sub,fill=f,line=l,tsize=11.5,ssize=8.5)
codebox(s,MX,TY+1820000,5300000,1150000,
        ['{"stage":"preview","tool":"issues_create",',
         ' "output_summary":{"approval_token_fp":',
         '   "ae7c28e1c18691be"}}',
         '// fingerprint — KHONG phai token song'],
        header="audit-2026-07-17.jsonl")
caption(s,"Bước 5 là bước duy nhất gọi APS ở lượt này — để kiểm issue_subtype_id có thật và đang active.")
setnotes(s,
 "Turn one. The LLM omits dry_run, so zod fills in true. The wrapper runs all six gates and stops at "
 "step six: it returns a full preview plus a token, and writes an audit entry. Note the log holds only "
 "a fingerprint of the token, never the live token.",
 "Lượt một. LLM không gửi dry_run nên zod điền true. Wrapper chạy hết sáu cổng rồi dừng ở bước sáu: trả "
 "preview đầy đủ kèm token, và ghi một dòng audit. Chú ý log chỉ giữ fingerprint của token, không bao giờ "
 "giữ token sống.")

# ===========================================================================
# 9 — Example, luot 2
# ===========================================================================
s=newslide("Ví dụ — lượt 2: thực thi",
           "Chạy lại TOÀN BỘ các cổng, rồi mới kiểm token — vì input có thể đã đổi giữa hai lượt.")
steps=[("1-5. chạy lại hết","allow-list, rate, rules, preview",PAPER,PAPER_L),
       ("6b. idempotency","key cũ + payload cũ → trả cache\nkhông gọi APS lần 2",MINT_F,GREEN),
       ("7. kiểm token","hết hạn / payload đổi → từ chối",AMBER_F,AMBER),
       ("8. EXECUTE","POST thật lên ACC",BLUE_F,BLUE_L),
       ("9. audit","stage: executed",MINT_F,GREEN)]
sh,w=hrow(s,steps,TY+300000,1350000,tsize=12.5,ssize=9)
connect(s,sh)
node(s,MX,TY+1900000,CW//2-130000,1000000,
     "Token bị XOÁ ngay khi dùng  (approval.ts:67)",
     "dùng lại lần hai → denied_approval, không phải failed_api",
     fill=RED_F,line=RED,tsize=13,ssize=10)
node(s,MX+CW//2+130000,TY+1900000,CW//2-130000,1000000,
     "APS trả 500  →  outcome_unknown",
     "POST đã bay đi rồi: ACC CÓ THỂ đã tạo issue trước khi lỗi",
     fill=AMBER_F,line=AMBER,tsize=13,ssize=10)
caption(s,"Ghi 'thất bại' khi không biết kết quả là nói sai sự thật cho người đọc log sau này.")
setnotes(s,
 "Turn two re-runs every gate before checking the token, because the input may have changed between "
 "calls. Two details worth remembering: the token is deleted the moment it is consumed, and a 500 on a "
 "mutation is recorded as outcome_unknown, not as a failure — the POST already left, so ACC may have "
 "applied it before erroring.",
 "Lượt hai chạy lại mọi cổng TRƯỚC khi kiểm token, vì input có thể đã đổi giữa hai lượt. Hai chi tiết đáng "
 "nhớ: token bị xoá ngay khi dùng; và lỗi 500 trên mutation được ghi là outcome_unknown chứ không phải thất "
 "bại — POST đã bay đi rồi nên ACC có thể đã áp dụng trước khi lỗi.")

# ===========================================================================
# 10 — What is a hash chain?
# ===========================================================================
s=newslide("Hash chain là gì?",
           "Mỗi dòng chứa hash của dòng trước — sửa một dòng giữa thì mọi dòng sau đều sai theo.")
y=TY+520000; bw=2500000; g=(CW-3*bw-700000)//2
xs=[MX+200000, MX+200000+bw+g, MX+200000+2*(bw+g)]
labs=[("Entry 1","prev: genesis\nthis: a3f...",PAPER,PAPER_L),
      ("Entry 2","prev: a3f...\nthis: 7c2...",PAPER,PAPER_L),
      ("Entry 3","prev: 7c2...\nthis: 91b...",PAPER,PAPER_L)]
shapes=[]
for i,(t,sub,f,l) in enumerate(labs):
    shapes.append(node(s,xs[i],y,bw,1100000,t,sub,fill=f,line=l,tsize=14,ssize=10))
connect(s,shapes)
codebox(s,MX,y+1450000,CW,1250000,
        ["this_hash = sha256( prev_hash + canonical_json(entry) )",
         "",
         "// canonical = SORT KEY truoc khi stringify (hash-chain.ts:6-10)",
         "// khong sort -> cung object, khac thu tu key -> khac hash"],
        header="src/safety/hash-chain.ts:4-13")
caption(s,"verifyChain kiểm HAI điều kiện — điều kiện thứ hai mới là cái bắt được XOÁ dòng.")
setnotes(s,
 "A hash chain makes a log tamper-evident: each line carries the hash of the previous one, so editing "
 "any earlier line invalidates every hash after it. The subtlety here is canonicalisation — keys are "
 "sorted before stringify, otherwise the same object in a different key order hashes differently. "
 "verifyChain also checks prev_hash adjacency, which is what catches a deleted line.",
 "Hash chain làm cho log phát hiện được gian lận: mỗi dòng mang hash của dòng trước, nên sửa bất kỳ dòng "
 "nào phía trước là mọi hash phía sau sai hết. Tinh tế ở chỗ canonical — phải sort key trước khi stringify, "
 "nếu không thì cùng một object mà khác thứ tự key sẽ ra hash khác. verifyChain còn kiểm prev_hash có trỏ "
 "đúng dòng trước không — đó mới là cái bắt được hành vi XOÁ nguyên một dòng.")

# ===========================================================================
# 11 — Audit trong project nay
# ===========================================================================
s=newslide("Cuốn sổ trong project này",
           "Mọi nhánh kết thúc đều ghi một dòng — kể cả các nhánh bị TỪ CHỐI, đó mới là phần có giá trị kiểm toán.")
r1=[("executed","đã làm thật",MINT_F,GREEN),
    ("preview","mới xem trước",BLUE_F,BLUE_L),
    ("idempotent_replay","trả cache, không gọi APS",MINT_F,GREEN),
    ("outcome_unknown","không biết kết quả",AMBER_F,AMBER)]
hrow(s,r1,TY+260000,780000,tsize=12,ssize=9)
r2=[("denied_allowlist",None,RED_F,RED),("denied_readonly",None,RED_F,RED),
    ("denied_rate_limit",None,RED_F,RED),("denied_auth_mode",None,RED_F,RED)]
hrow(s,r2,TY+1180000,560000,tsize=11.5,ssize=9)
r3=[("denied_approval",None,RED_F,RED),("denied_missing_approval",None,RED_F,RED),
    ("denied_business_rule",None,RED_F,RED),("failed_api",None,RED_F,RED)]
hrow(s,r3,TY+1840000,560000,tsize=11.5,ssize=9)
node(s,MX,TY+2600000,CW,800000,
     "Ranh giới đã khai báo trung thực: call sai schema bị MCP SDK chặn TRƯỚC handler → không có dòng nào",
     "manifest không nói 'audit every tool call' mà nói 'every invocation that reaches a handler'",
     fill=PAPER,line=PAPER_L,tsize=13,ssize=10)
caption(s,"Audit fail-open mặc định; FORMA_AUDIT_FAIL_CLOSED=true để huỷ cả lệnh khi không ghi được log.")
setnotes(s,
 "Twelve stages. The denied_* ones matter most for an audit: they record what was attempted and refused. "
 "One honest boundary is declared rather than hidden: a call whose input fails the SDK schema check never "
 "reaches a handler, so it produces no entry — nothing was done on that path.",
 "Mười hai stage. Các stage denied_* mới là phần quan trọng khi kiểm toán: chúng ghi lại điều đã bị thử và "
 "bị từ chối. Một ranh giới được khai báo trung thực thay vì giấu: call sai schema bị SDK chặn trước handler "
 "nên không sinh dòng nào — vì đường đó không làm gì cả.")

# ===========================================================================
# 12 — Allow-list: bon loai scope
# ===========================================================================
s=newslide("Allow-list: mỗi tool tự khai mình thuộc loại nào",
           "Allow-list chứa DM id (b.<guid>) — nhưng không phải tool nào cũng nhận id loại đó.")
items=[("dm","input có DM hub/project id\n→ CHECK",BLUE_F,BLUE_L),
       ("discovery","không có input để scope\n→ tự LỌC output",MINT_F,GREEN),
       ("no-resource","không đụng resource ACC\n(meta_*) → bỏ qua",PAPER,PAPER_L),
       ("unmappable","id không map về DM được\n→ TỪ CHỐI",RED_F,RED)]
sh,_=hrow(s,items,TY+380000,1400000,tsize=15,ssize=10)
node(s,MX,TY+2050000,CW,780000,
     "aecdm_* , md_* , docs_get_viewables , issues_pin_element  =  unmappable",
     "AECDM có id-space riêng; /modelderivative/{urn} không scope theo project → không có gì để đối chiếu",
     fill=PAPER,line=PAPER_L,tsize=13,ssize=10)
node(s,MX,TY+2960000,CW,560000,
     "Mặc định FORMA_ALLOWED_HUBS=* và _PROJECTS=*  →  allow-list TẮT  →  không tool nào bị chặn",
     None,fill=MINT_F,line=GREEN,tsize=12.5)
caption(s,"Enforce tại src/tools/_wrap.ts:95-118 (enforceScope). scope là field bắt buộc: _types.ts:30-43.")
setnotes(s,
 "The allow-list holds Data Management ids, but not every tool takes one. So each tool declares its case. "
 "The refusal branch is the honest answer for AECDM and Model Derivative tools: their ids cannot be mapped "
 "back to a DM project, so letting them through would break the promise the manifest makes.",
 "Allow-list chứa DM id, nhưng không phải tool nào cũng nhận id loại đó. Nên mỗi tool tự khai trường hợp của "
 "mình. Nhánh từ chối là câu trả lời trung thực cho AECDM và Model Derivative: id của chúng không map ngược "
 "về DM project được, cho qua là phá vỡ lời hứa trong manifest.")

# ===========================================================================
# 13 — Vi sao khai bao thay vi doan
# ===========================================================================
s=newslide("Vì sao phải KHAI BÁO scope thay vì đoán từ tên field?",
           "Vì đoán từ tên field đã tạo ra lỗ hổng thật — đây là bug lớn nhất từng có trong repo.")
node(s,MX,TY+240000,CW,620000,
     "Wrapper cũ: đọc bất kỳ input nào TÊN LÀ project_id / hub_id",
     None,fill=PAPER,line=PAPER_L,tsize=14)
bad=[("Sót hoàn toàn","5 tool AECDM dùng element_group_id\nmd_* dùng urn  →  KHÔNG HỀ bị kiểm",RED_F,RED),
     ("Kiểm NHẦM giá trị","aecdm_list_element_groups có field TÊN project_id\nnhưng chứa AECDM id  →  so với DM id",RED_F,RED)]
hrow(s,bad,TY+1000000,1300000,tsize=13.5,ssize=10)
node(s,MX,TY+2480000,CW,560000,
     "Kết quả: allow-list bật vẫn đọc được mọi model  +  từ chối oan các call hợp lệ",
     None,fill=RED_F,line=RED,tsize=13)
node(s,MX,TY+3160000,CW,620000,
     "Bài học: định danh phải theo NAMESPACE, không theo TÊN BIẾN",
     "khai báo tường minh + kiểu bắt buộc  →  quên = lỗi biên dịch, không phải lỗ hổng im lặng",
     fill=MINT_F,line=GREEN,tsize=13.5,ssize=10)
caption(s,"tests/unit/tools/registry-scope.spec.ts khoá thêm các bất biến mà compiler không diễn đạt được.")
setnotes(s,
 "The old wrapper picked the id to check by field name. Two consequences: tools scoped by another id were "
 "never checked at all, and one tool whose project_id field holds an AECDM id had it compared against DM "
 "ids — protecting nothing while rejecting valid calls. Identity must follow namespace, not variable name.",
 "Wrapper cũ chọn id để kiểm theo TÊN field. Hai hậu quả: tool dùng id tên khác thì không hề bị kiểm, và một "
 "tool có field project_id chứa AECDM id lại bị đem so với DM id — không bảo vệ gì mà còn từ chối oan. "
 "Định danh phải theo namespace, không theo tên biến.")

# ===========================================================================
# 14 — SSA vs 2LO
# ===========================================================================
s=newslide("Hai danh tính, KHÔNG phải failover",
           "Ở mode ssa server tạo cả hai token; mỗi tool chọn một lần lúc startup — 2LO lỗi không thử lại bằng SSA.")
node(s,MX+200000,TY+320000,4700000,1300000,"SSA  —  danh tính 'robot'",
     "được mời vào project • thấy đúng project được gán\nBẮT BUỘC cho Issues / Reviews / AECDM / MC",
     fill=BLUE_F,line=BLUE_L,tsize=15,ssize=10)
node(s,RIGHT-4900000,TY+320000,4700000,1300000,"2LO  —  danh tính 'ứng dụng'",
     "client_credentials • thấy MỌI project trong hub\ndùng cho DM / Admin / MD / Docs (preferredAuth)",
     fill=AMBER_F,line=AMBER,tsize=15,ssize=10)
node(s,MX,TY+1900000,CW,780000,
     "Manifest từng khai '2LO ... with SSA as fallback'  —  fallback CHƯA TỪNG TỒN TẠI",
     "index.ts:44 LUÔN gán auth2lo ở mode ssa → nhánh kiểm ctx.auth2lo không bao giờ rơi về SSA",
     fill=RED_F,line=RED,tsize=13.5,ssize=10)
node(s,MX,TY+2800000,CW,720000,
     "Bài học: claim thừa kế KHÔNG phải bằng chứng",
     "câu đó nghe hợp lý nên sống sót nhiều vòng review — đến khi có người đối chiếu với code",
     fill=MINT_F,line=GREEN,tsize=13.5,ssize=10)
caption(s,"Cùng họ với lỗi app_model: ghi chú nội bộ nói 'FAQ không định nghĩa A/B/C' — định nghĩa nằm TRANG 1.")
setnotes(s,
 "Two identities with different visibility. In ssa mode the server builds both, and each tool picks one at "
 "startup. This is not a failover: a failed 2LO request is never retried under SSA. The manifest claimed a "
 "fallback that never existed — a plausible sentence survived several review rounds because nobody checked "
 "it against the code.",
 "Hai danh tính với tầm nhìn khác nhau. Ở mode ssa server tạo cả hai, mỗi tool chọn một lúc startup. Đây "
 "KHÔNG phải failover: 2LO lỗi thì không bao giờ thử lại bằng SSA. Manifest từng khai một fallback chưa từng "
 "tồn tại — câu nghe hợp lý nên sống sót qua nhiều vòng review vì không ai đối chiếu với code.")

# ===========================================================================
# 15 — Bay da tra gia
# ===========================================================================
s=newslide("Những cái bẫy đã trả giá thật",
           "Bốn cái bẫy này không phải lý thuyết — mỗi cái đều tốn ít nhất một vòng review để phát hiện.")
t1=[("Verifier tự chế cho GREEN GIẢ","regex báo '0 vi phạm' trong khi còn 6\n193 test xanh mà exe sai version",RED_F,RED),
    ("Im lặng bị đọc thành PASS","probe in ALLOWED! cho 16 tool\nthực ra server chết lúc khởi động",RED_F,RED)]
hrow(s,t1,TY+280000,1350000,tsize=13,ssize=9.5)
t2=[("Một sự thật, NĂM bản sao","runtime • manifest • form • README • email\nsửa 3 sót 2, lặp lại 5 vòng",AMBER_F,AMBER),
    ("hookTimeout ≠ testTimeout","nâng testTimeout 30s nhưng suite import\ntrong beforeAll vẫn bị chặn ở 10s",AMBER_F,AMBER)]
hrow(s,t2,TY+1780000,1350000,tsize=13,ssize=9.5)
node(s,MX,TY+3280000,CW,620000,
     "Cách duy nhất dứt vòng lặp: biến check tay thành check máy",
     "manifest-sync.spec.ts  •  version-sync.spec.ts  •  registry-scope.spec.ts  •  env-free.spec.ts",
     fill=MINT_F,line=GREEN,tsize=13.5,ssize=10)
caption(s,"Probe phải phân biệt BA trạng thái: đạt / vi phạm / không phản hồi — im lặng không bao giờ là pass.")
setnotes(s,
 "Four traps that each cost a review round. Home-made verifiers give false greens; silence gets read as a "
 "pass; one fact duplicated in five places drifts; and raising testTimeout does not help a suite that "
 "imports inside beforeAll. The only thing that ended the loop was turning hand checks into machine checks.",
 "Bốn cái bẫy, mỗi cái tốn một vòng review. Verifier tự chế cho green giả; im lặng bị đọc thành pass; một sự "
 "thật nhân bản ở năm nơi thì sẽ lệch; và nâng testTimeout không cứu được suite import trong beforeAll. Thứ "
 "duy nhất chấm dứt vòng lặp là biến check tay thành check máy.")

# ===========================================================================
# 16 — AECDM vs MD
# ===========================================================================
s=newslide("Vì sao AECDM không có 'Level'?",
           "Giới hạn của NGUỒN DỮ LIỆU, không phải do query chưa đủ thông minh — và có guardrail chặn LLM đi vào ngõ cụt.")
node(s,MX+200000,TY+300000,4700000,1250000,"AECDM  (GraphQL)",
     "chỉ phơi ra VALUE parameter\ngiữ Base Offset, Elevation at Bottom\nMẤT Level / Base Constraint / Host",
     fill=RED_F,line=RED,tsize=15,ssize=10)
node(s,RIGHT-4900000,TY+300000,4700000,1250000,"Model Derivative  (SVF2)",
     "phơi ra TOÀN BỘ parameter Revit\nmd_get_properties + group_by\n→ tổng diện tích theo tầng",
     fill=MINT_F,line=GREEN,tsize=15,ssize=10)
node(s,MX,TY+1820000,CW,700000,
     "Level / Base Constraint / Host là REFERENCE parameter — trỏ tới element khác, không phải giá trị",
     "AECDM làm phẳng tập property nên chỉ giữ được các scalar, mất luôn cái trỏ tới node Level",
     fill=PAPER,line=PAPER_L,tsize=13,ssize=10)
node(s,MX,TY+2650000,CW,760000,
     "Guardrail runtime: src/tools/aecdm/_param-guidance.ts",
     "LLM gọi aggregate với group_by='Level' → short-circuit + chỉ đường sang md_get_properties",
     fill=BLUE_F,line=BLUE_L,tsize=13.5,ssize=10)
caption(s,"Live-verified: Floors 203 across 14 levels, tổng 72,671 ft²; Walls 3046 across 16 groups.")
setnotes(s,
 "A data-source limit, not a smartness gap. Level and Base Constraint are Revit reference parameters — "
 "links to another element — and AECDM's flattened property set keeps only value parameters. So per-storey "
 "grouping must go through Model Derivative. A runtime guardrail redirects the LLM instead of letting it "
 "dead-end.",
 "Giới hạn của nguồn dữ liệu, không phải do kém thông minh. Level và Base Constraint là reference parameter "
 "của Revit — trỏ tới element khác — mà AECDM làm phẳng chỉ giữ value parameter. Nên muốn nhóm theo tầng "
 "phải qua Model Derivative. Có guardrail runtime chỉ đường cho LLM thay vì để nó đi vào ngõ cụt.")

# ===========================================================================
# 17 — Takeaways
# ===========================================================================
s=newslide("5 điều phải nhớ",
           "Nếu chỉ nhớ năm điều từ cả codebase này, nhớ năm điều dưới đây.")
items=[("1","An toàn phải là MẶC ĐỊNH","dry_run=true tự tiêm vào mọi mutation tool — ai quên đọc docs vẫn an toàn"),
       ("2","Dồn an toàn vào MỘT chỗ","46 tool × 8 lớp = 368 chỗ có thể quên; _wrap.ts lo hết thì tool thứ 47 tự động được bảo vệ"),
       ("3","Khai báo > suy đoán","scope là field bắt buộc → quên là lỗi biên dịch, không phải lỗ hổng im lặng"),
       ("4","Tuyên bố phải khớp HÀNH VI","manifest/README/PRIVACY là artifact bị review — nói 'audit mọi thứ' mà fail-open là nói dối"),
       ("5","Check tay sẽ lệch, check máy thì không","5 vòng review chỉ dứt khi viết manifest-sync.spec.ts")]
y=TY+180000; h=640000; g=130000
for i,(n,t,sub) in enumerate(items):
    yy=y+i*(h+g)
    circle(s,MX,yy+h//2-190000,380000,n,fill=BLUE_L)
    node(s,MX+520000,yy,CW-520000,h,t,sub,
         fill=MINT_F if i in (0,4) else PAPER, line=GREEN if i in (0,4) else PAPER_L,
         tsize=14,ssize=10,align='l')
caption(s,"Chi tiết đầy đủ + câu hỏi tự kiểm: docs/study/STUDY_toan-bo-codebase.md")
setnotes(s,
 "Five things to carry away. Safety must be the default, not an option. Concentrate it in one place so new "
 "tools inherit it. Prefer declaration over inference so mistakes become compile errors. Keep declarations "
 "matching behaviour, because they are reviewed artifacts. And turn hand checks into machine checks — that "
 "is what finally ended five rounds of review.",
 "Năm điều mang về. An toàn phải là mặc định chứ không phải tuỳ chọn. Dồn nó vào một chỗ để tool mới tự động "
 "thừa hưởng. Ưu tiên khai báo hơn suy đoán để lỗi trở thành lỗi biên dịch. Giữ tuyên bố khớp hành vi, vì đó "
 "là artifact bị review. Và biến check tay thành check máy — đó là thứ chấm dứt năm vòng review.")

# ---------------------------------------------------------------------------
os.makedirs(_HERE, exist_ok=True)
prs.save(OUT)
print("OK ->", OUT)
print("slides:", len(prs.slides._sldIdLst))
